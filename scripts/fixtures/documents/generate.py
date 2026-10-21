"""Regenerate synthetic fixtures; requires python-docx, openpyxl, python-pptx, reportlab, pypdf.

These documents contain no third-party or personal content. Kept as binary fixtures
so normal Node tests do not require Python, Microsoft Office or LibreOffice.
"""
from pathlib import Path
from io import BytesIO
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches as PptxInches
from reportlab.pdfgen.canvas import Canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.lib.utils import ImageReader
from pypdf import PdfReader, PdfWriter
from PIL import Image, ImageDraw

out = Path(__file__).parent
document = Document()
image = Image.new('RGB', (400, 120), '#e4efe9')
draw = ImageDraw.Draw(image)
draw.rectangle((24, 24, 72, 96), fill='#277756')
draw.rectangle((96, 48, 144, 96), fill='#4f8d70')
draw.rectangle((168, 12, 216, 96), fill='#71a38b')
draw.text((244, 48), 'HEMI 1500', fill='#1f4935')
image_bytes = BytesIO()
image.save(image_bytes, format='PNG')

document.add_heading('真实 Word 中文测试', 0)
document.add_paragraph('项目预算与交付计划')
table = document.add_table(rows=2, cols=2)
table.style = 'Light Shading Accent 1'
for cell, text in zip([table.cell(0, 0), table.cell(0, 1), table.cell(1, 0), table.cell(1, 1)], ['项目', '金额', '软件', '1500']):
    cell.text = text
document.add_picture(BytesIO(image_bytes.getvalue()), width=Inches(4))
paragraph = document.add_paragraph()
link = OxmlElement('w:hyperlink')
link.set(qn('r:id'), paragraph.part.relate_to('javascript:parent.__HEMI_DOC_LINK_EXECUTED__=true', RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
run = OxmlElement('w:r')
text = OxmlElement('w:t')
text.text = '链接安全检查'
run.append(text)
link.append(run)
paragraph._p.append(link)
document.save(out / 'chinese.docx')

workbook = Workbook()
sheet = workbook.active
sheet.title = '预算'
sheet.append(['项目', '金额'])
sheet.append(['软件', 1500])
workbook.save(out / 'chinese.xlsx')

presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[1])
slide.shapes.title.text = '真实 PPT 中文测试'
slide.placeholders[1].text = '项目预算 1500'
slide.shapes.add_picture(BytesIO(image_bytes.getvalue()), PptxInches(1), PptxInches(3), width=PptxInches(5))
second = presentation.slides.add_slide(presentation.slide_layouts[1])
second.shapes.title.text = '第二张：结论'
second.placeholders[1].text = '按期交付，预算 1500'
presentation.save(out / 'chinese.pptx')

pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
canvas = Canvas(str(out / 'chinese.pdf'))
canvas.setFont('STSong-Light', 14)
canvas.drawString(72, 720, '真实 PDF 中文测试：预算 1500')
canvas.drawImage(ImageReader(image), 72, 540, width=300, height=90)
canvas.save()

scan = Canvas(str(out / 'scan.pdf'))
scan.drawImage(ImageReader(image), 72, 540, width=400, height=120)
scan.save()

encrypted = PdfWriter()
for page in PdfReader(out / 'chinese.pdf').pages:
    encrypted.add_page(page)
encrypted.encrypt('synthetic-test-password')
with open(out / 'encrypted.pdf', 'wb') as output:
    encrypted.write(output)
